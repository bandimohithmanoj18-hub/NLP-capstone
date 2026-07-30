import os
import re
from pathlib import Path
from typing import Optional, Tuple
from sqlalchemy.orm import Session
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT

from app.models.complaint import Complaint
from app.core.config import settings
from app.core.logger import get_logger

logger = get_logger("app.services.document_service")


class DocumentService:
    """
    Formal Legal Document Generator producing OOXML (.docx) and ReportLab (.pdf) deliverables.
    Part of **Milestone 8: PDF/DOCX generation**.
    """

    @staticmethod
    def _sanitize_filename(name: str) -> str:
        """Sanitizes party name to build a safe filename without path traversal/invalid chars."""
        # Replace spaces with underscores
        name_clean = name.replace(' ', '_')
        # Retain only letters, numbers, underscores, and hyphens
        return re.sub(r'[^a-zA-Z0-9_\-]', '', name_clean)

    @staticmethod
    def generate_complaint_documents(db: Session, complaint_id: int) -> Tuple[Path, Path]:
        """Generates both .docx and .pdf complaint files and updates database paths."""
        complaint = db.query(Complaint).filter(Complaint.id == complaint_id).first()
        if not complaint:
            raise ValueError(f"Complaint ID {complaint_id} not found.")

        settings.ensure_directories()
        clean_name = DocumentService._sanitize_filename(complaint.complainant_name)
        base_name = f"Complaint_{complaint.id}_{clean_name}"
        docx_path = settings.GENERATED_DIR / f"{base_name}.docx"
        pdf_path = settings.GENERATED_DIR / f"{base_name}.pdf"

        # Generate modern OOXML DOCX
        DocumentService._create_docx(complaint, docx_path)

        # Generate high-resolution ReportLab PDF
        DocumentService._create_pdf(complaint, pdf_path)

        complaint.generated_docx_path = str(docx_path)
        complaint.generated_pdf_path = str(pdf_path)
        complaint.status = "READY"
        db.add(complaint)
        db.commit()
        db.refresh(complaint)

        logger.info(f"Generated DOCX & PDF for complaint ID {complaint_id}")
        return docx_path, pdf_path

    @staticmethod
    def generate_legal_notice(db: Session, complaint_id: int) -> Tuple[Path, Path]:
        """Generates a Pre-Litigation Formal Legal Notice addressed to the Opposite Party."""
        complaint = db.query(Complaint).filter(Complaint.id == complaint_id).first()
        if not complaint:
            raise ValueError(f"Complaint ID {complaint_id} not found.")

        settings.ensure_directories()
        clean_name = DocumentService._sanitize_filename(complaint.opposite_party_name)
        base_name = f"Legal_Notice_{complaint.id}_{clean_name}"
        docx_path = settings.GENERATED_DIR / f"{base_name}.docx"
        pdf_path = settings.GENERATED_DIR / f"{base_name}.pdf"

        # Generate Notice DOCX
        DocumentService._create_notice_docx(complaint, docx_path)

        # Generate Notice PDF
        DocumentService._create_notice_pdf(complaint, pdf_path)

        logger.info(f"Generated Pre-Litigation Legal Notice for complaint ID {complaint_id}")
        return docx_path, pdf_path

    @staticmethod
    def _create_docx(complaint: Complaint, output_path: Path) -> None:
        doc = docx.Document()

        # Page Setup
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.2)
            section.right_margin = Inches(1.0)

        # Dynamic State Extraction from address
        state = "STATE OF COMPLAINANT"
        addr_upper = (complaint.complainant_address or "").upper()
        for st in [
            "ANDHRA PRADESH", "ARUNACHAL PRADESH", "ASSAM", "BIHAR", "CHHATTISGARH", "GOA", "GUJARAT",
            "HARYANA", "HIMACHAL PRADESH", "JHARKHAND", "KARNATAKA", "KERALA", "MADHYA PRADESH",
            "MAHARASHTRA", "MANIPUR", "MEGHALAYA", "MIZORAM", "NAGALAND", "ODISHA", "PUNJAB",
            "RAJASTHAN", "SIKKIM", "TAMIL NADU", "TELANGANA", "TRIPURA", "UTTAR PRADESH",
            "UTTARAKHAND", "WEST BENGAL", "DELHI", "JAMMU & KASHMIR", "LADAKH", "PUDUCHERRY"
        ]:
            if st in addr_upper:
                state = st
                break

        # Title Heading
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_para.add_run(
            f"BEFORE THE {complaint.jurisdiction_forum.replace('_', ' ').upper()} AT {state}\n"
            f"CONSUMER COMPLAINT NO. _____ OF 2026"
        )
        run.bold = True
        run.font.size = Pt(14)

        doc.add_paragraph("IN THE MATTER OF:").paragraph_format.space_after = Pt(6)

        # Complainant block
        p_comp = doc.add_paragraph()
        run_comp = p_comp.add_run(f"{complaint.complainant_name}\n")
        run_comp.bold = True
        p_comp.add_run(f"{complaint.complainant_address}\nContact: {complaint.complainant_contact}\n")
        p_comp.add_run("... COMPLAINANT\n\nVERSUS\n\n").bold = True

        # Opposite Party block
        p_op = doc.add_paragraph()
        run_op = p_op.add_run(f"{complaint.opposite_party_name}\n")
        run_op.bold = True
        p_op.add_run(f"{complaint.opposite_party_address}\nContact: {complaint.opposite_party_contact}\n")
        p_op.add_run("... OPPOSITE PARTY").bold = True

        p_subject = doc.add_paragraph()
        run_subject = p_subject.add_run(f"SUBJECT: COMPLAINT UNDER SECTION 35 OF THE CONSUMER PROTECTION ACT, 2019 FOR DISPUTED CLAIM OF ₹{complaint.claim_amount:,.2f}")
        run_subject.bold = True

        # Section 1: Facts
        doc.add_heading("I. FACTS OF THE CASE", level=2)
        for line in (complaint.facts or "").split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())

        # Section 2: Grounds
        doc.add_heading("II. GROUNDS OF RELIEF", level=2)
        for line in (complaint.grounds or "").split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())

        # Section 3: Relief Sought
        doc.add_heading("III. RELIEF SOUGHT / PRAYER", level=2)
        for line in (complaint.relief_sought or "").split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())

        # Section 4: Verification
        doc.add_heading("IV. VERIFICATION", level=2)
        doc.add_paragraph(complaint.verification_clause or "")
        doc.add_paragraph("\n\n___________________________\nCOMPLAINANT SIGNATURE\nDate: _______________")

        doc.save(str(output_path))

    @staticmethod
    def _create_pdf(complaint: Complaint, output_path: Path) -> None:
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=54,
        )
        styles = getSampleStyleSheet()
        normal = styles["Normal"]
        normal.fontSize = 10
        normal.leading = 14

        # Dynamic State Extraction from address
        state = "STATE OF COMPLAINANT"
        addr_upper = (complaint.complainant_address or "").upper()
        for st in [
            "ANDHRA PRADESH", "ARUNACHAL PRADESH", "ASSAM", "BIHAR", "CHHATTISGARH", "GOA", "GUJARAT",
            "HARYANA", "HIMACHAL PRADESH", "JHARKHAND", "KARNATAKA", "KERALA", "MADHYA PRADESH",
            "MAHARASHTRA", "MANIPUR", "MEGHALAYA", "MIZORAM", "NAGALAND", "ODISHA", "PUNJAB",
            "RAJASTHAN", "SIKKIM", "TAMIL NADU", "TELANGANA", "TRIPURA", "UTTAR PRADESH",
            "UTTARAKHAND", "WEST BENGAL", "DELHI", "JAMMU & KASHMIR", "LADAKH", "PUDUCHERRY"
        ]:
            if st in addr_upper:
                state = st
                break

        title_style = ParagraphStyle(
            "LegalTitle",
            parent=normal,
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=16,
            alignment=TA_CENTER,
        )
        h2_style = ParagraphStyle(
            "LegalH2",
            parent=normal,
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=15,
            spaceBefore=10,
            spaceAfter=6,
        )

        story = []
        story.append(Paragraph(
            f"BEFORE THE {complaint.jurisdiction_forum.replace('_', ' ').upper()} AT {state}<br/>"
            f"<b>CONSUMER COMPLAINT NO. _____ OF 2026</b>",
            title_style
        ))
        story.append(Spacer(1, 15))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#2563EB"), spaceAfter=12))

        # Parties table
        party_data = [
            [Paragraph(f"<b>{complaint.complainant_name}</b><br/>{complaint.complainant_address}<br/>Tel: {complaint.complainant_contact}", normal),
             Paragraph("<b>... COMPLAINANT</b>", normal)],
            [Paragraph("<b>VERSUS</b>", normal), ""],
            [Paragraph(f"<b>{complaint.opposite_party_name}</b><br/>{complaint.opposite_party_address}<br/>Tel: {complaint.opposite_party_contact}", normal),
             Paragraph("<b>... OPPOSITE PARTY</b>", normal)],
        ]
        table = Table(party_data, colWidths=[350, 130])
        table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ]))
        story.append(table)
        story.append(Spacer(1, 12))

        story.append(Paragraph(f"<b>SUBJECT: COMPLAINT UNDER SECTION 35 OF CPA 2019 FOR CLAIM OF ₹{complaint.claim_amount:,.2f}</b>", normal))
        story.append(Spacer(1, 10))

        story.append(Paragraph("I. FACTS OF THE CASE", h2_style))
        for line in (complaint.facts or "").split("\n"):
            if line.strip():
                story.append(Paragraph(line.strip(), normal))
                story.append(Spacer(1, 4))

        story.append(Paragraph("II. GROUNDS OF RELIEF", h2_style))
        for line in (complaint.grounds or "").split("\n"):
            if line.strip():
                story.append(Paragraph(line.strip(), normal))
                story.append(Spacer(1, 4))

        story.append(Paragraph("III. RELIEF SOUGHT / PRAYER", h2_style))
        for line in (complaint.relief_sought or "").split("\n"):
            if line.strip():
                story.append(Paragraph(line.strip(), normal))
                story.append(Spacer(1, 4))

        story.append(Paragraph("IV. VERIFICATION CLAUSE", h2_style))
        story.append(Paragraph(complaint.verification_clause or "", normal))
        story.append(Spacer(1, 20))
        story.append(Paragraph("<b>____________________________</b><br/><b>COMPLAINANT SIGNATURE</b><br/>Date: _______________", normal))

        doc.build(story)

    @staticmethod
    def _create_notice_docx(complaint: Complaint, output_path: Path) -> None:
        doc = docx.Document()
        doc.add_heading("PRE-LITIGATION LEGAL NOTICE", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"To,\n{complaint.opposite_party_name}\n{complaint.opposite_party_address}\nEmail/Tel: {complaint.opposite_party_contact}\n").runs[0].bold = True
        doc.add_paragraph(f"SUBJECT: STATUTORY LEGAL NOTICE UNDER SECTION 35 & 47 OF CONSUMER PROTECTION ACT, 2019 FOR DISPUTED CLAIM OF ₹{complaint.claim_amount:,.2f}").runs[0].bold = True

        doc.add_paragraph(
            f"Under instructions from and on behalf of my client / the Complainant ({complaint.complainant_name}), "
            f"you are hereby served with the present Legal Notice:\n\n"
            f"1. That you have committed a Deficiency in Service and Unfair Trade Practice by failing to resolve the genuine grievance of the Complainant.\n"
            f"2. You are hereby called upon to refund/pay the disputed amount of ₹{complaint.claim_amount:,.2f} within 15 days of receiving this notice.\n"
            f"3. In the event of your failure to comply within 15 days, my client shall be constrained to initiate formal proceedings before the {complaint.jurisdiction_forum.replace('_', ' ').title()} without further notice, holding you liable for all costs and consequences."
        )
        doc.add_paragraph("\n\n___________________________\nLEGAL COUNSEL / ADVOCATE FOR COMPLAINANT\nDate: _______________")
        doc.save(str(output_path))

    @staticmethod
    def _create_notice_pdf(complaint: Complaint, output_path: Path) -> None:
        doc = SimpleDocTemplate(str(output_path), pagesize=A4, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
        styles = getSampleStyleSheet()
        normal = styles["Normal"]
        title_style = ParagraphStyle("NoticeTitle", parent=normal, fontName="Helvetica-Bold", fontSize=13, leading=16, alignment=TA_CENTER)

        story = [
            Paragraph("<b>PRE-LITIGATION LEGAL NOTICE</b>", title_style),
            Spacer(1, 15),
            HRFlowable(width="100%", thickness=1, color=colors.HexColor("#7E22CE"), spaceAfter=12),
            Paragraph(f"<b>TO:</b><br/><b>{complaint.opposite_party_name}</b><br/>{complaint.opposite_party_address}<br/>Tel/Email: {complaint.opposite_party_contact}", normal),
            Spacer(1, 12),
            Paragraph(f"<b>SUBJECT: LEGAL NOTICE FOR DEFICIENCY IN SERVICE (CLAIM: ₹{complaint.claim_amount:,.2f})</b>", normal),
            Spacer(1, 12),
            Paragraph(
                f"Under instructions from and on behalf of my client ({complaint.complainant_name}), you are hereby called upon to refund "
                f"the disputed claim of ₹{complaint.claim_amount:,.2f} within <b>15 days</b> of receipt of this notice, failing which legal proceedings "
                f"shall be instituted before the {complaint.jurisdiction_forum.replace('_', ' ').title()} under CPA 2019.",
                normal
            ),
            Spacer(1, 30),
            Paragraph("<b>____________________________</b><br/><b>ADVOCATE / LEGAL COUNSEL</b><br/>Date: _______________", normal),
        ]
        doc.build(story)
